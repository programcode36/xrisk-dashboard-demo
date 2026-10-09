#!/usr/bin/env python3
"""
Descarga el reporte "Tasa de Interés Promedio del Sistema Bancario"
(tasas activas por tipo de crédito y por empresa bancaria) desde el sitio
PUBLICO de la SBS:
https://www.sbs.gob.pe/app/pp/EstadisticasSAEEPortal/Paginas/TIActivaTipoCreditoEmpresa.aspx?tip=B

A diferencia de los reportes B-2201/B-2401 (archivos .XLS fijos en la
intranet), esta es una pagina web dinamica tipo ASP.NET (Telerik
RadControls), protegida por Incapsula (firewall/anti-bots).

Historial de intentos (para quien edite esto despues):
1. Peticion directa con `requests`: Incapsula devuelve una pagina de
   bloqueo de ~800 bytes (status 200, pero "Incapsula incident...") en
   vez del contenido real. Confirmado con una prueba real.
2. Abrir Chrome con Selenium solo para "resolver" el reto y despues pasar
   las cookies a una sesion de `requests`: Incapsula lo sigue bloqueando.
   Probablemente porque tambien valida el fingerprint TLS/HTTP del
   cliente (la libreria `requests` no "se ve" como Chrome aunque tenga
   las cookies correctas), no solo la cookie de sesion.
3. SOLUCION ACTUAL: hacer TODA la interaccion (poner la fecha, elegir
   moneda, hacer clic en "Exportar") dentro del propio navegador real
   controlado por Selenium, configurando Chrome para que descargue el
   archivo a una carpeta fija sin preguntar. Como nunca se sale del
   navegador real, no hay mismatch de fingerprint posible.

Elementos relevantes de la pagina (confirmados contra el HTML real que
compartio el usuario):
    ctl00_cphContent_rdpDate_dateInput  -> input visible de fecha (DD/MM/YYYY)
    ctl00_cphContent_lbtnMn             -> pestaña "Moneda Nacional"
    ctl00_cphContent_lbtnMex            -> pestaña "Moneda Extranjera"
    ctl00_cphContent_btnExportar        -> boton "Exportar"
Los botones/pestañas estan dentro de un UpdatePanel ("updConsulta"), asi
que cambiar de pestaña hace un postback asincrono (AJAX) que puede
reemplazar esos nodos del DOM; por eso cada elemento se vuelve a buscar
por su id justo antes de usarlo, en vez de guardar la referencia de mas
arriba (evita errores de "elemento obsoleto").

Requisitos: Chrome instalado (Selenium >= 4.6 descarga el chromedriver
compatible automaticamente, no hace falta instalarlo a mano) y
`pip install selenium`.

Uso:
    python descarga_sbs_tasas_activas.py --fecha 07/10/2026
    python descarga_sbs_tasas_activas.py --fecha 07/10/2026 --moneda ME
    python descarga_sbs_tasas_activas.py --fecha 07/10/2026 --moneda ambas
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from pathlib import Path

import pandas as pd

URL = "https://www.sbs.gob.pe/app/pp/EstadisticasSAEEPortal/Paginas/TIActivaTipoCreditoEmpresa.aspx?tip=B"

PROCESSED_DIR = Path("data/processed")

ID_FECHA = "ctl00_cphContent_rdpDate_dateInput"
ID_TAB_MN = "ctl00_cphContent_lbtnMn"
ID_TAB_ME = "ctl00_cphContent_lbtnMex"
ID_BOTON_EXPORTAR = "ctl00_cphContent_btnExportar"


def _crear_driver(carpeta_descargas: Path):
    from selenium import webdriver

    carpeta_descargas.mkdir(parents=True, exist_ok=True)
    opciones = webdriver.ChromeOptions()
    opciones.add_argument("--window-size=1280,900")
    # Sin --headless a proposito: Incapsula distingue Chrome headless del
    # normal con mas facilidad.
    opciones.add_experimental_option(
        "prefs",
        {
            "download.default_directory": str(carpeta_descargas.resolve()),
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
        },
    )
    return webdriver.Chrome(options=opciones)


def _esperar_elemento(driver, id_elemento: str, timeout: int = 30):
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC

    return WebDriverWait(driver, timeout).until(EC.presence_of_element_located((By.ID, id_elemento)))


def _poner_fecha(driver, fecha: str):
    from selenium.webdriver.common.keys import Keys

    campo = _esperar_elemento(driver, ID_FECHA)
    campo.click()
    campo.send_keys(Keys.CONTROL, "a")
    campo.send_keys(Keys.DELETE)
    campo.send_keys(fecha)
    campo.send_keys(Keys.TAB)  # dispara el blur para que el control sincronice la fecha
    time.sleep(1)


def _elegir_moneda(driver, moneda: str):
    if moneda == "MN":
        tab = _esperar_elemento(driver, ID_TAB_MN)
    else:
        tab = _esperar_elemento(driver, ID_TAB_ME)
    tab.click()
    time.sleep(2)  # la pestaña hace un postback asincrono (AJAX) dentro del UpdatePanel


def _esperar_descarga(carpeta: Path, archivos_antes: set, timeout: int = 30) -> Path:
    fin = time.time() + timeout
    while time.time() < fin:
        actuales = set(carpeta.iterdir())
        nuevos = [
            p for p in (actuales - archivos_antes)
            if p.is_file() and not p.name.endswith(".crdownload") and not p.name.endswith(".tmp")
        ]
        if nuevos:
            # Esperar un instante mas por si el archivo sigue escribiendose
            time.sleep(1)
            return nuevos[0]
        time.sleep(0.5)
    raise TimeoutError(f"No aparecio ningun archivo nuevo en {carpeta} despues de {timeout}s")


def descargar_tasas(driver, carpeta_descargas: Path, fecha: str, moneda: str) -> Path:
    from selenium.webdriver.common.by import By

    driver.get(URL)
    _esperar_elemento(driver, ID_BOTON_EXPORTAR)

    _elegir_moneda(driver, moneda)
    _poner_fecha(driver, fecha)

    archivos_antes = set(carpeta_descargas.iterdir()) if carpeta_descargas.exists() else set()

    boton = driver.find_element(By.ID, ID_BOTON_EXPORTAR)
    boton.click()

    return _esperar_descarga(carpeta_descargas, archivos_antes)


def limpiar_tabla(path_crudo: Path, moneda: str) -> pd.DataFrame:
    """
    Limpia el archivo tal como lo exporta la pagina de la SBS (titulo,
    subtitulo, columna en blanco a la izquierda, y nota al pie) y se
    queda solo con la tabla: encabezado "Tasa Anual (%)" + bancos, filas
    de datos, con una columna "Moneda" agregada justo despues de la
    columna de tipo de credito. Verificado que coincide, valor por valor,
    con una limpieza manual de referencia.
    """
    df = pd.read_excel(path_crudo, sheet_name=0, header=None)
    df = df.loc[:, ~df.isna().all(axis=0)]  # quita la columna A, que viene siempre en blanco
    df = df.reset_index(drop=True)

    fila_header = None
    for r in range(len(df)):
        fila = df.iloc[r].astype(str).str.strip()
        if (fila == "Tasa Anual (%)").any():
            fila_header = r
            break
    if fila_header is None:
        raise ValueError(f"No se encontro la fila de encabezado 'Tasa Anual (%)' en {path_crudo}")

    encabezados = df.iloc[fila_header].tolist()

    fila_fin = len(df)
    for r in range(fila_header + 1, len(df)):
        if pd.isna(df.iloc[r, 0]):  # primera fila en blanco despues de los datos = fin de la tabla
            fila_fin = r
            break

    tabla = df.iloc[fila_header + 1 : fila_fin].reset_index(drop=True)
    tabla.columns = encabezados
    tabla.insert(1, "Moneda", moneda)
    return tabla


def _es_tipo_credito(texto) -> bool:
    """
    Distingue una fila de "tipo de credito" (ej. 'Corporativos', 'Grandes
    Empresas', 'Consumo') de sus sub-items (ej. 'Descuentos', 'Prestamos
    hasta 30 dias'): en el archivo de la SBS, los tipos de credito vienen
    con 8 o mas espacios de sangria al inicio del texto, y los sub-items
    con 5. Verificado contra el archivo real (9-10 espacios en los tipos
    de credito, 5 en los sub-items, para las 7 categorias del reporte).
    """
    if not isinstance(texto, str):
        return False
    return (len(texto) - len(texto.lstrip(" "))) >= 8


def guardar_excel(df: pd.DataFrame, fecha, destino: Path) -> None:
    """
    Guarda el consolidado con una columna "Fecha" agregada al inicio
    (formateada dd/mm/aaaa), encabezado en negrita, y en negrita solo la
    celda de "Tasa Anual (%)" de las filas que son tipo de credito (no
    sus sub-items).
    """
    df = df.copy()
    col_tipo_credito = df.columns[0]  # "Tasa Anual (%)", antes de insertar Fecha
    es_categoria = df[col_tipo_credito].map(_es_tipo_credito)
    df.insert(0, "Fecha", fecha)

    df.to_excel(destino, sheet_name="Reporte", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(destino)
    ws = wb["Reporte"]

    negrita = Font(bold=True)
    for celda in ws[1]:
        celda.font = negrita

    col_fecha = 1
    col_tipo = 2  # "Fecha" queda en la A, "Tasa Anual (%)" pasa a la B
    for fila_excel in range(2, ws.max_row + 1):
        ws.cell(row=fila_excel, column=col_fecha).number_format = "dd/mm/yyyy"
        if es_categoria.iloc[fila_excel - 2]:
            ws.cell(row=fila_excel, column=col_tipo).font = negrita

    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 32
    ws.freeze_panes = "A2"
    wb.save(destino)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--fecha", required=True, metavar="DD/MM/YYYY", help="Fecha a consultar, ej. 07/10/2026")
    ap.add_argument(
        "--moneda",
        choices=["MN", "ME", "ambas"],
        default="ambas",
        help="MN = solo Moneda Nacional, ME = solo Moneda Extranjera, ambas = las dos juntas en un solo Excel (default)",
    )
    ap.add_argument("--outdir", default=str(PROCESSED_DIR), help="Carpeta de salida final")
    args = ap.parse_args()

    if not re.match(r"^\d{2}/\d{2}/\d{4}$", args.fecha):
        print("La fecha debe tener el formato DD/MM/YYYY, ej. 07/10/2026", file=sys.stderr)
        sys.exit(1)

    salida = Path(args.outdir)
    salida.mkdir(parents=True, exist_ok=True)
    carpeta_descargas = salida / "_descargas_temp"

    monedas = ["MN", "ME"] if args.moneda == "ambas" else [args.moneda]

    print("Abriendo Chrome (no lo cierres hasta que termine)...")
    try:
        driver = _crear_driver(carpeta_descargas)
    except Exception as exc:
        print(f"[ERROR] No se pudo abrir Chrome con Selenium: {exc}", file=sys.stderr)
        print("Verifica que Chrome este instalado y que 'pip install selenium' se haya hecho bien.", file=sys.stderr)
        sys.exit(1)

    tablas = []
    try:
        for moneda in monedas:
            print(f"Exportando tasas activas ({moneda}) al {args.fecha}...")
            archivo_crudo = None
            try:
                archivo_crudo = descargar_tasas(driver, carpeta_descargas, args.fecha, moneda)
                tabla = limpiar_tabla(archivo_crudo, moneda)
            except Exception as exc:
                print(f"  [ERROR] {exc}", file=sys.stderr)
                continue
            finally:
                if archivo_crudo is not None and archivo_crudo.exists():
                    archivo_crudo.unlink()
            print(f"  OK: {len(tabla)} filas de {moneda}")
            tablas.append(tabla)
    finally:
        driver.quit()
        if carpeta_descargas.exists() and not any(carpeta_descargas.iterdir()):
            carpeta_descargas.rmdir()

    if not tablas:
        print("No se logro descargar ninguna tabla.", file=sys.stderr)
        sys.exit(1)

    consolidado = pd.concat(tablas, ignore_index=True)
    dia, mes, anio = args.fecha.split("/")
    fecha_obj = pd.Timestamp(year=int(anio), month=int(mes), day=int(dia)).date()
    destino = salida / f"tasas_activas_bancos_{anio}-{mes}-{dia}.xlsx"
    guardar_excel(consolidado, fecha_obj, destino)
    print(f"\nListo: {len(consolidado)} filas guardadas en {destino}")


if __name__ == "__main__":
    main()
